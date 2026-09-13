import io, re, pdfplumber, openpyxl
import streamlit as st
from openpyxl.styles import Font, PatternFill, Alignment
from collections import defaultdict

# --- Streamlit Web Interface Setup ---
st.set_page_config(
    page_title="Shipment Tag Reviewer",
    page_icon="📦",
    layout="centered"
)

st.title("📦 Shipment Tag & Order Reviewer")
st.write("Upload your PDF airwaybill and Excel order sheet below to generate the reconciliation report.")

pdf_file = st.file_uploader("1. Upload Airwaybill PDF", type=["pdf"])
xlsx_file = st.file_uploader("2. Upload Ops Excel Sheet (.xlsx)", type=["xlsx"])

if st.button("🚀 Run Review", type="primary"):
    if not pdf_file or not xlsx_file:
        st.error("Please upload both the PDF and Excel file before running.")
    else:
        with st.spinner("Processing files... Please wait."):
            # ---------- Extract tag data from PDF ----------
            amt_re = re.compile(r'\u0645\.\u062c(\d[\d,\.]*)')
            ref_re = re.compile(r'(\d{6,12})\s*Order Reference:\s*trimize:#(\d+)')
            item_re = re.compile(r'x\s*(\d+)\s*\(([A-Z0-9\-]+)\)')

            tags = []
            with pdfplumber.open(pdf_file) as pdf:
                for page_num, page in enumerate(pdf.pages, start=1):
                    text = page.extract_text() or ""
                    ref = ref_re.search(text)
                    amt = amt_re.search(text)
                    items = item_re.findall(text)
                    tags.append({
                        "page": page_num,
                        "tracking_number": ref.group(1) if ref else None,
                        "order_id": int(ref.group(2)) if ref else None,
                        "tag_amount": float(amt.group(1).replace(",", "")) if amt else None,
                        "tag_items": [(sku, int(qty)) for qty, sku in items],
                    })

            # ---------- Load & group Excel sheet ----------
            wb_in = openpyxl.load_workbook(xlsx_file, data_only=True)
            ws_in = wb_in["Ops"]
            headers = [c.value for c in ws_in[1]]
            idx = {h: i for i, h in enumerate(headers)}
            rows = list(ws_in.iter_rows(min_row=2, values_only=True))

            orders = defaultdict(lambda: {"skus": []})
            for r in rows:
                oid = r[idx["Name"]]
                if oid is None:
                    continue
                o = orders[oid]
                if r[idx["Total"]] not in (None, ""):
                    o["total"] = r[idx["Total"]]
                    o["payment_method"] = r[idx["Payment Method"]]
                    o["shipment_id"] = r[idx["Shipment ID"]]
                    o["financial_status"] = r[idx["Financial Status"]]
                    o["outstanding"] = r[idx["Outstanding Balance"]]
                if r[idx["Lineitem sku"]]:
                    o["skus"].append((r[idx["Lineitem sku"]], r[idx["Lineitem quantity"]]))

            def payment_type(pm):
                if not pm:
                    return "Unknown"
                return "Visa/Online" if "Paymob" in pm else "Cash (COD)"

            def approx_equal(a, b, tol=0.01):
                if a is None or b is None:
                    return a == b
                return abs(float(a) - float(b)) <= tol

            # ---------- Compare tags against order ----------
            results = []
            for t in tags:
                oid = t["order_id"]
                order = orders.get(oid)

                row = {
                    "Page": t["page"],
                    "Order ID (tag)": oid,
                    "Tracking # (tag)": t["tracking_number"],
                    "Tag Amount (COD)": t["tag_amount"] if t["tag_amount"] is not None else "None (prepaid)",
                    "Tag Items (SKU x Qty)": ", ".join(f"{s} x{q}" for s, q in t["tag_items"]),
                }

                if order is None:
                    row.update({
                        "Found in Orders Sheet": "NO - order not in sheet",
                        "Excel Shipment ID": "", "Shipment ID Match": "",
                        "Excel Amount Due": "", "Amount Match": "",
                        "Tag Payment Type": "Cash (COD)" if t["tag_amount"] is not None else "Visa/Online",
                        "Excel Payment Method": "", "Payment Match": "",
                        "Excel Items (SKU x Qty)": "", "Items Match": "",
                        "Overall Flag": "ORDER NOT FOUND",
                    })
                    results.append(row)
                    continue

                excel_shipment = order.get("shipment_id")
                shipment_match = (str(excel_shipment) == str(t["tracking_number"]))

                excel_due = order.get("outstanding")
                tag_amt = t["tag_amount"]
                if tag_amt is None:
                    amount_match = approx_equal(excel_due, 0)
                else:
                    amount_match = approx_equal(excel_due, tag_amt)

                tag_pay_type = "Cash (COD)" if tag_amt is not None else "Visa/Online"
                excel_pay_type = payment_type(order.get("payment_method"))
                payment_match = (tag_pay_type == excel_pay_type)

                excel_skus = sorted(order.get("skus", []))
                tag_skus = sorted(t["tag_items"])
                items_match = (excel_skus == tag_skus)

                flags = []
                if not shipment_match: flags.append("Shipment ID")
                if not amount_match: flags.append("Amount")
                if not payment_match: flags.append("Payment")
                if not items_match: flags.append("Items")
                overall = "OK" if not flags else "CHECK: " + ", ".join(flags)

                row.update({
                    "Found in Orders Sheet": "Yes",
                    "Excel Shipment ID": excel_shipment,
                    "Shipment ID Match": "Yes" if shipment_match else "NO",
                    "Excel Amount Due": excel_due,
                    "Amount Match": "Yes" if amount_match else "NO",
                    "Tag Payment Type": tag_pay_type,
                    "Excel Payment Method": order.get("payment_method"),
                    "Payment Match": "Yes" if payment_match else "NO",
                    "Excel Items (SKU x Qty)": ", ".join(f"{s} x{q}" for s, q in excel_skus),
                    "Items Match": "Yes" if items_match else "NO",
                    "Overall Flag": overall,
                })
                results.append(row)

            # ---------- Write output workbook to memory ----------
            cols = list(results[0].keys())
            wb_out = openpyxl.Workbook()
            ws = wb_out.active
            ws.title = "Tag Review"

            header_font = Font(name="Arial", bold=True, color="FFFFFF")
            header_fill = PatternFill("solid", fgColor="4472C4")
            ok_fill = PatternFill("solid", fgColor="C6E0B4")
            bad_fill = PatternFill("solid", fgColor="F8CBAD")

            for c, col_name in enumerate(cols, start=1):
                cell = ws.cell(row=1, column=c, value=col_name)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", wrap_text=True)

            for r, row in enumerate(results, start=2):
                for c, col_name in enumerate(cols, start=1):
                    cell = ws.cell(row=r, column=c, value=row[col_name])
                    cell.font = Font(name="Arial")
                flag_cell = ws.cell(row=r, column=cols.index("Overall Flag") + 1)
                flag_cell.fill = ok_fill if row["Overall Flag"] == "OK" else bad_fill

            widths = {"Page":6, "Order ID (tag)":14, "Tracking # (tag)":16, "Tag Amount (COD)":16,
                      "Tag Items (SKU x Qty)":26, "Found in Orders Sheet":18, "Excel Shipment ID":16,
                      "Shipment ID Match":14, "Excel Amount Due":14, "Amount Match":12,
                      "Tag Payment Type":14, "Excel Payment Method":34, "Payment Match":12,
                      "Excel Items (SKU x Qty)":26, "Items Match":12, "Overall Flag":26}
            for c, col_name in enumerate(cols, start=1):
                ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = widths.get(col_name, 14)
            ws.freeze_panes = "A2"

            buffer = io.BytesIO()
            wb_out.save(buffer)
            buffer.seek(0)

            n_ok = sum(1 for r in results if r["Overall Flag"] == "OK")
            n_flag = len(results) - n_ok

            st.success("✅ Finished processing!")
            st.metric(label="Total Tags Processed", value=len(results))
            st.write(f"**OK:** {n_ok} | **Flagged:** {n_flag}")

            st.download_button(
                label="📥 Download Shipment_Tag_Review.xlsx",
                data=buffer,
                file_name="Shipment_Tag_Review.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
